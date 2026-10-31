"""Проверки прототипа извлечения TXT/DOCX (пуш 4, Р4)."""

from __future__ import annotations

import base64
import io

import pytest
from docx import Document

from anongw.contracts import UnsupportedInputError
from anongw.files.extractor import TxtDocxFileExtractor

pytestmark = pytest.mark.unit


def _b64(data: bytes) -> str:
    return base64.b64encode(data).decode("ascii")


def _docx_bytes() -> bytes:
    document = Document()
    document.add_paragraph("Справка выдана Ивану Петрову.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Телефон"
    table.rows[0].cells[1].text = "+7 900 000-00-00"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


async def test_txt_utf8() -> None:
    extractor = TxtDocxFileExtractor()
    text = await extractor.extract_text(_b64("Привет, Иван Петров".encode()), "letter.txt")
    assert text == "Привет, Иван Петров"


async def test_txt_utf16_with_bom() -> None:
    extractor = TxtDocxFileExtractor()
    raw = "Отправлено Анной Смирновой".encode("utf-16")
    text = await extractor.extract_text(_b64(raw), "note.txt")
    assert text == "Отправлено Анной Смирновой"


async def test_txt_cp1251_fallback() -> None:
    extractor = TxtDocxFileExtractor()
    raw = "Адрес: г. Тверь, ул. Мира, д. 7".encode("cp1251")
    text = await extractor.extract_text(_b64(raw), "addr.txt")
    assert text == "Адрес: г. Тверь, ул. Мира, д. 7"


async def test_docx_paragraphs_and_table() -> None:
    extractor = TxtDocxFileExtractor()
    text = await extractor.extract_text(_b64(_docx_bytes()), "cert.docx")
    assert "Ивану Петрову" in text
    assert "Телефон; +7 900 000-00-00" in text


async def test_unsupported_format_rejected_before_model() -> None:
    extractor = TxtDocxFileExtractor()
    with pytest.raises(UnsupportedInputError, match="TXT и DOCX"):
        await extractor.extract_text(_b64(b"%PDF-1.4 ..."), "scan.pdf")


async def test_filename_with_path_rejected() -> None:
    extractor = TxtDocxFileExtractor()
    with pytest.raises(UnsupportedInputError, match="путь"):
        await extractor.extract_text(_b64("текст".encode()), "../secret/letter.txt")
