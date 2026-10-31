"""Извлечение текста из TXT/DOCX (прототип недели 1, пуш 4).

Наружу отдаётся только текст: имя файла, метаданные и встроенные объекты
не возвращаются и не пересылаются внешней модели (ТЗ п.2).
Поддерживаемые части DOCX и ограничения — docs/files/README.md.
"""

from __future__ import annotations

import base64
import binascii
import io

from anongw.config import Settings
from anongw.contracts import UnsupportedInputError

_ALLOWED_EXTENSIONS = frozenset({".txt", ".docx"})


def _validate_filename(filename: str | None) -> str:
    """Возвращает расширение; неподдерживаемое — отказ ДО обращения к модели."""
    if filename is None or not filename.strip():
        raise UnsupportedInputError("Имя файла обязательно: по нему определяется формат")
    name = filename.strip()
    if "/" in name or "\\" in name:
        raise UnsupportedInputError("Имя файла не должно содержать путь")
    dot = name.rfind(".")
    ext = name[dot:].lower() if dot != -1 else ""
    if ext not in _ALLOWED_EXTENSIONS:
        shown = ext if ext else "без расширения"
        raise UnsupportedInputError(
            f"Формат «{shown}» не поддерживается: разрешены только TXT и DOCX"
        )
    return ext


def _decode_txt(data: bytes) -> str:
    """TXT: UTF-8 (с BOM и без), UTF-16 LE/BE по BOM, запасной вариант CP1251."""
    if data.startswith(b"\xef\xbb\xbf"):
        return data.decode("utf-8-sig")
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16")
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        return data.decode("cp1251")


def _extract_docx(data: bytes) -> str:
    """DOCX: параграфы и таблицы; метаданные документа не читаем вовсе."""
    from docx import Document  # импорт внутри: зависимость нужна только для DOCX

    try:
        document = Document(io.BytesIO(data))
        parts = [p.text for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    parts.append("; ".join(cells))
    except Exception as exc:  # повреждённые/не-DOCX файлы дают разные исключения библиотеки
        raise UnsupportedInputError(f"Не удалось прочитать DOCX: {exc}") from exc
    return "\n".join(parts).strip()


class TxtDocxFileExtractor:
    """Разбирает inline file_data: base64 → байты → текст нужного формата."""

    def __init__(self, settings: Settings | None = None) -> None:
        max_mb = settings.max_file_mb if settings is not None else 10
        self._max_bytes = max_mb * 1024 * 1024

    async def extract_text(self, file_data_b64: str, filename: str | None) -> str:
        try:
            data = base64.b64decode(file_data_b64, validate=True)
        except (binascii.Error, ValueError) as exc:
            raise UnsupportedInputError("file_data не является корректным base64") from exc

        if len(data) > self._max_bytes:
            limit_mb = self._max_bytes // (1024 * 1024)
            raise UnsupportedInputError(f"Файл больше ограничения {limit_mb} МБ")

        ext = _validate_filename(filename)
        if ext == ".txt":
            return _decode_txt(data)

        text = _extract_docx(data)
        if not text:
            raise UnsupportedInputError("Не удалось извлечь текст из документа")
        return text
